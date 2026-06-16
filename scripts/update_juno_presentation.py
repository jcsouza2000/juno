#!/usr/bin/env python3
"""Atualiza Juno_Presentation.pptx com marcos juno_audit v0.3 e piloto 2026-06."""
from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
PPT = ROOT / "Juno_Presentation.pptx"
BACKUP = ROOT / "Juno_Presentation.pptx.bak"
VALIDATION_DATE = "2026-06-13"
VERSION = "v0.3.2"
TOTAL_SLIDES = 12


def _replace_exact(slide, old: str, new: str) -> int:
    n = 0
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip() == old:
            shape.text = new
            n += 1
    return n


def _replace_contains(slide, replacements: dict[str, str]) -> None:
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        txt = shape.text
        for old, new in replacements.items():
            if old in txt:
                shape.text = txt.replace(old, new)


def _update_footers(prs: Presentation) -> None:
    for slide in prs.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            t = shape.text
            if "JUNO Industrial Diagnostic" not in t or "produto validado" not in t:
                continue
            if "|" in t and "/" in t.split("|")[-1]:
                idx = t.split("|")[-1].strip().split("/")[0].zfill(2)
                shape.text = (
                    f"JUNO Industrial Diagnostic {VERSION} | "
                    f"produto validado em {VALIDATION_DATE} | {idx}/{TOTAL_SLIDES}"
                )
            elif "2026-05-25" in t:
                shape.text = t.replace("2026-05-25", VALIDATION_DATE)


def main() -> None:
    if not PPT.exists():
        raise SystemExit(f"Arquivo nao encontrado: {PPT}")

    if not BACKUP.exists():
        shutil.copy2(PPT, BACKUP)

    prs = Presentation(str(PPT))

    # Slide 1
    _replace_contains(prs.slides[0], {
        "Produto validado tecnicamente em C:\\Souza\\juno": (
            "Piloto validado: E2E 10/10 + juno-verify TRILHA INTEGRA (prod local Docker)"
        ),
        "IA local governada | auditoria | conectores ERP | relatorios financeiros": (
            "IA governada | juno_audit (ledger + OTS) | ERP | relatorios executivos"
        ),
    })

    # Slide 5
    _replace_contains(prs.slides[4], {
        "action confirmation antes de executar mudancas": (
            "action confirmation + trilha juno_audit em cada /score"
        ),
        "Resultado: IA deixa de ser chatbot e vira uma interface operacional governada.": (
            "Resultado: IA operacional com trilha verificavel — export bundle + juno-verify ao vivo."
        ),
    })

    # Slide 9
    _replace_contains(prs.slides[8], {
        "audit_logs com usuario, recurso, mudanca, contexto e diff": (
            "audit_logs + juno_audit (cadeia, Merkle, OTS, Agent Registry v31)"
        ),
    })

    # Slide 10 — prova tecnica
    s10 = prs.slides[9]
    _replace_contains(s10, {
        "A apresentacao do produto se apoia em uma base validada, nao em promessa.": (
            "Base validada em prod local Docker — ledger, ancoragem OTS e export verificavel."
        ),
    })
    _replace_exact(s10, "190", "262")
    _replace_exact(s10, "testes backend passando", "testes backend (pytest)")
    _replace_exact(s10, "vulnerabilidades npm moderadas+", "smoke E2E piloto API")
    _replace_exact(s10, "erros mypy", "juno-verify TRILHA INTEGRA")
    _replace_exact(s10, "190 passed", "262 passed")
    _replace_exact(s10, "sem warnings residuais", "backend + juno_audit no Docker")
    _replace_exact(s10, "Postgres, Redis, backend e frontend", "Postgres :8002 / frontend :4002")
    _replace_exact(s10, "mypy / ruff / black", "juno_audit suites")
    _replace_exact(s10, "tipagem e estilo estruturados", "smoke + anchor + sprint3 + router")
    _replace_exact(s10, "pytest -q --no-cov", "pytest backend")
    _replace_exact(s10, "Docker real", "Docker prod local")
    # dois blocos "0" — primeiro -> 10/10, segundo -> OK
    zeros = [sh for sh in s10.shapes if hasattr(sh, "text") and sh.text.strip() == "0"]
    if len(zeros) >= 1:
        zeros[0].text = "10/10"
    if len(zeros) >= 2:
        zeros[1].text = "OK"

    # Slide 11
    _replace_contains(prs.slides[10], {
        "um produto com testes, build, auditoria e persistencia": (
            "produto com testes, ledger verificavel (juno-verify) e persistencia Postgres"
        ),
    })

    # Slide 12
    _replace_contains(prs.slides[11], {
        "Validar login, fluxo /ai, upload financeiro/PDF e um conector ERP "
        "com dados de demonstracao ou cliente piloto.": (
            "Demo ao vivo: registrar agente, /score, fechar lote (SUBMITTED), "
            "baixar bundle, juno-verify TRILHA INTEGRA — depois ERP/financeiro."
        ),
        "rodar diagnostico inicial": "rodar /score + scheduler (--min-batch 1)",
        "validar acoes governadas": "validar GET /v31/audit/export/{tenant}",
        "preparar pacote executivo": "salvar print juno-verify para pagina /trust",
        "politica de retencao e auditoria": "JUNO_AGENT_SK + migrations v1-v3",
    })

    _update_footers(prs)

    prs.save(str(PPT))
    print(f"Atualizado: {PPT}")
    print(f"Backup: {BACKUP}")


if __name__ == "__main__":
    main()
