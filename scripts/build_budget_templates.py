"""
Gera templates de Orçamento (PT) e Budget (EN) a partir dos templates existentes.

1. Adiciona colunas Orçamento e Budget ao workbook canonico Templates/Demontra*.xlsx
2. Cria template_demonstracoes_orcamento.xlsx e template_demonstracoes_budget.xlsx
   (formato longo conta/periodo/valor para upload em Demonstracoes)
"""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES_DIR = ROOT / "Templates"
UPLOAD_SRC = ROOT / "janus" / "frontend" / "public" / "templates" / "template_demonstracoes.xlsx"
UPLOAD_ORC = ROOT / "janus" / "frontend" / "public" / "templates" / "template_demonstracoes_orcamento.xlsx"
UPLOAD_BUD = ROOT / "janus" / "frontend" / "public" / "templates" / "template_demonstracoes_budget.xlsx"
BACKEND_ORC = ROOT / "janus" / "backend" / "template_demonstracoes_orcamento.xlsx"
BACKEND_BUD = ROOT / "janus" / "backend" / "template_demonstracoes_budget.xlsx"

# Metas anuais 2025 (R$ milhoes) — realizado ~20.697 mi; orcamento ligeiramente acima (+~3%)
BUDGET_2025_MIL = {
    "Receita Líquida": 21_350.0,
    "EBITDA": 8_050.0,
    "EBITDA excl. efeitos não recorrentes": 8_050.0,
    "Lucro (prejuízo) Líquido": 1_720.0,
    "EBITDA Ajustado": 8_050.0,
    "EBITDA Ajustado Excluídos Efeitos Não Recorrentes": 8_050.0,
}

QUARTER_KEYS = ("1T25", "2T25", "3T25", "4T25")
REALIZED_QUARTERS = {
    "Receita Líquida": (4858.534, 5247.202, 5426.467, 5165.304),
    "EBITDA": (1858.811, 2040.855, 2116.94149942, 1831.5091368700023),
    "EBITDA excl. efeitos não recorrentes": (1858.811, 2040.855, 2116.94149942, 1831.5091368700023),
}


def _find_canonical() -> Path:
    matches = sorted(p for p in TEMPLATES_DIR.glob("Demontra*.xlsx") if not p.name.startswith("~$"))
    if not matches:
        raise FileNotFoundError(f"Nenhum Demontra*.xlsx em {TEMPLATES_DIR}")
    return matches[0]


def _header_col(ws, header_row: int, label: str) -> int | None:
    for col in range(1, ws.max_column + 1):
        val = ws.cell(header_row, col).value
        if val is None:
            continue
        text = str(val).strip().replace(".0", "")
        if text.lower() == label.lower() or text == label:
            return col
    return None


def _insert_budget_columns(wb_path: Path) -> None:
    wb = openpyxl.load_workbook(wb_path)
    for sheet_name in ("2. DRE", "5. EBITDA"):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        col_2025 = _header_col(ws, 1, "2025")
        if col_2025 is None:
            continue
        insert_at = col_2025 + 1
        ws.insert_cols(insert_at, amount=2)
        ws.cell(1, insert_at, "Orçamento")
        ws.cell(1, insert_at + 1, "Budget")

        annual_key = None
        for row in range(2, ws.max_row + 1):
            label = str(ws.cell(row, 1).value or "").strip()
            if not label or label.lower() == "nan":
                continue
            annual = BUDGET_2025_MIL.get(label)
            if annual is not None:
                ws.cell(row, insert_at, annual)
                ws.cell(row, insert_at + 1, annual)
                annual_key = label
            elif label in REALIZED_QUARTERS:
                realized = REALIZED_QUARTERS[label]
                total_r = sum(realized)
                total_b = BUDGET_2025_MIL.get(label, BUDGET_2025_MIL.get("Receita Líquida", total_r))
                if label != "Receita Líquida":
                    total_b = BUDGET_2025_MIL.get(label, total_r * (21_350 / 20_697.507))
                factor = total_b / total_r if total_r else 1.0
                for q_idx, q_key in enumerate(QUARTER_KEYS):
                    q_col = _header_col(ws, 1, q_key)
                    if q_col and q_idx < len(realized):
                        val = round(realized[q_idx] * factor, 3)
                        ws.cell(row, q_col, val)

        if annual_key:
            print(f"  {sheet_name}: colunas Orçamento/Budget inseridas após 2025")

    wb.save(wb_path)
    print(f"Atualizado: {wb_path.name}")


def _build_upload_budget(src: Path, dest: Path, period_label: str, scale: float, readme: str) -> None:
    shutil.copy2(src, dest)
    wb = openpyxl.load_workbook(dest)
    if "Leia-me" in wb.sheetnames:
        del wb["Leia-me"]
    readme_ws = wb.create_sheet("Leia-me", 0)
    readme_ws["A1"] = readme
    readme_ws.column_dimensions["A"].width = 90

    for sheet in ("DRE", "Balanco", "DFC"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for row in range(2, ws.max_row + 1):
            period_cell = ws.cell(row, 2)
            val_cell = ws.cell(row, 3)
            if period_cell.value is None:
                continue
            period_cell.value = period_label
            if isinstance(val_cell.value, (int, float)):
                val_cell.value = round(float(val_cell.value) * scale, 2)
    wb.save(dest)
    print(f"Criado: {dest.relative_to(ROOT)}")


def main() -> None:
    canonical = _find_canonical()
    backup = canonical.with_suffix(".xlsx.bak")
    if not backup.exists():
        shutil.copy2(canonical, backup)
        print(f"Backup: {backup.name}")

    _insert_budget_columns(canonical)

    if not UPLOAD_SRC.exists():
        raise FileNotFoundError(UPLOAD_SRC)

    _build_upload_budget(
        UPLOAD_SRC,
        UPLOAD_ORC,
        "2025-ORC",
        1.05,
        "Template de ORÇAMENTO 2025 (Demonstrações). "
        "Preencha e importe em Demonstrações; use a coluna Orçamento do template "
        "executivo Templates/ para comparativos por cenário (Fase C). "
        "periodo=2025-ORC identifica linhas de meta orçamentária.",
    )
    _build_upload_budget(
        UPLOAD_SRC,
        UPLOAD_BUD,
        "2025-BUD",
        1.05,
        "BUDGET template 2025 (Financial Statements). "
        "Fill and import under Demonstrações; the executive Templates workbook "
        "uses the Budget column for scenario comparatives (Phase C). "
        "period=2025-BUD marks budget target rows.",
    )

    shutil.copy2(UPLOAD_ORC, BACKEND_ORC)
    shutil.copy2(UPLOAD_BUD, BACKEND_BUD)

    # Validacao rapida
    dre = pd.read_excel(canonical, sheet_name="2. DRE", header=None)
    headers = [str(x).strip() for x in dre.iloc[0].tolist()]
    assert "Orçamento" in headers and "Budget" in headers
    idx = headers.index("Orçamento")
    for i, row in dre.iterrows():
        if str(row.iloc[0]).strip() == "Receita Líquida":
            print(f"Validacao: Orçamento Receita Líquida = {row.iloc[idx]}")
            break


if __name__ == "__main__":
    main()
